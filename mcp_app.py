import argparse
import asyncio
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any
from urllib import error, request

from document_reader import DocumentReader
from mcp.server.mcpserver.server import MCPServer
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph


WORD_PATTERN = re.compile(r"[^\W_]+(?:['’-][^\W_]+)*", re.UNICODE)
PLACEHOLDER_PATTERN = re.compile(r"<[A-Za-z][A-Za-z0-9_ -]*>")
URL_PATTERN = re.compile(r"https?://[^\s<>\"\\]+")
SHELL_COMMAND_PATTERN = re.compile(r"^(?:ansible-playbook|cd|pwd|unzip)\b")
STRICT_GENERATION_SYSTEM_PROMPT = """You are a strict document-fidelity assistant.
Use only facts and words found in the supplied source document. Do not add,
infer, paraphrase, summarize, or omit information. Preserve the requested
Markdown or YAML structure. Return only converted content."""
STRICT_VALIDATION_SYSTEM_PROMPT = """You are a strict document-fidelity verifier.
Return exactly VALID only when the converted output preserves every source fact
and adds no words or facts. Otherwise return exactly INVALID."""


def _markdown_runs(paragraph: Paragraph) -> str:
    """Preserve the inline formatting Markdown can represent."""
    def wrap(text: str, marker: str) -> str:
        """Keep whitespace outside emphasis markers so CommonMark can parse them."""
        leading = text[: len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()):]
        content = text.strip()
        return f"{leading}{marker}{content}{marker}{trailing}" if content else text

    def format_run(run: Any) -> str:
        text = run.text.replace("\n", "<br>")
        if not text:
            return ""
        if run.bold and run.italic:
            text = wrap(text, "***")
        elif run.bold:
            text = wrap(text, "**")
        elif run.italic:
            text = wrap(text, "*")
        if run.font.strike:
            text = wrap(text, "~~")
        # CommonMark has no portable underline syntax. Emit the text itself
        # instead of inserting HTML <u> tags into Markdown output.
        return text

    parts: list[str] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            label = "".join(format_run(run) for run in item.runs).strip()
            url = item.url or item.address
            parts.append(f"[{label}]({url})" if label and url else label)
        else:
            parts.append(format_run(item))
    rendered = "".join(parts).replace("****", "").strip() or paragraph.text.strip()
    return PLACEHOLDER_PATTERN.sub(lambda match: f"`{match.group(0)}`", rendered)


def _table_cell_text(cell: Any) -> str:
    text = "<br>".join(_markdown_runs(paragraph) for paragraph in cell.paragraphs if paragraph.text.strip())
    # Word split this URL across paragraphs; a Markdown line-break marker makes
    # it invalid when serialized to YAML.
    if "http" in text.casefold():
        url_match = URL_PATTERN.search(text.replace("<br>", ""))
        return url_match.group(0) if url_match else text.replace("<br>", "")
    return text


def _table_cell_source_text(cell: Any) -> str:
    """Keep raw Word cell text for strict YAML conversion."""
    return "\n".join(paragraph.text for paragraph in cell.paragraphs if paragraph.text)


def _is_list_paragraph(paragraph: Paragraph, style: str) -> bool:
    paragraph_properties = paragraph._p.pPr
    return (
        style.startswith("list bullet")
        or style.startswith("list numbering")
        or (
            not any(run.bold for run in paragraph.runs)
            and paragraph_properties is not None
            and paragraph_properties.numPr is not None
        )
    )


def _code_language(text: str) -> str | None:
    stripped = text.strip()
    if SHELL_COMMAND_PATTERN.match(stripped):
        return "bash"
    if re.match(r"^(?:def |class |import |from |print\()", stripped):
        return "python"
    if (stripped.startswith("{") and stripped.endswith("}")) or (stripped.startswith("[") and stripped.endswith("]")):
        try:
            json.loads(stripped)
            return "json"
        except json.JSONDecodeError:
            return None
    if re.match(r"^[A-Za-z_][\w-]*:\s*\S+", stripped):
        return "yaml"
    return None


def _word_list_number(paragraph: Paragraph, counters: dict[str, list[int]]) -> str | None:
    """Reconstruct Word automatic numbering, which python-docx omits from text."""
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is None or paragraph_properties.numPr is None:
        return None
    number_properties = paragraph_properties.numPr
    number_id = str(number_properties.numId.val)
    level = int(number_properties.ilvl.val)
    values = counters.setdefault(number_id, [])
    while len(values) <= level:
        values.append(0)
    values[level] += 1
    del values[level + 1:]
    return ".".join(str(value) for value in values)


def extract_blocks(doc: Any) -> list[dict[str, Any]]:
    """Extract blocks in their original Word document order."""
    blocks: list[dict[str, Any]] = []
    list_counters: dict[str, list[int]] = {}
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, doc)
            plain_text = paragraph.text.strip()
            text = _markdown_runs(paragraph)
            automatic_number = _word_list_number(paragraph, list_counters)
            if not text:
                continue

            style = (paragraph.style.name or "").lower()
            if style.startswith("heading") or style == "title":
                # Reserve # for the DOCX title; Word Heading 1 begins at ##.
                level = 1 if style == "title" else int(style.split()[-1]) + 1
                block = {"type": "heading", "level": min(level, 6), "text": text, "source_text": plain_text}
                if automatic_number:
                    block["automatic_number"] = automatic_number
                blocks.append(block)
            elif _is_list_paragraph(paragraph, style):
                list_type = "ordered" if "number" in style or re.match(r"^\d+[.)]", plain_text) else "bullet"
                block: dict[str, Any] = {"type": list_type, "text": text, "source_text": plain_text}
                if list_type == "ordered":
                    number_match = re.match(r"^(\d+(?:\.\d+)*[.)]?)\s*", plain_text)
                    text = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", text)
                    block["text"] = text
                    if number_match:
                        block["number"] = number_match.group(1).rstrip(".)")
                blocks.append(block)
            elif set(plain_text) == {"-"} and len(plain_text) >= 3:
                blocks.append({"type": "rule"})
            elif language := _code_language(plain_text):
                blocks.append({"type": "code", "language": language, "text": text, "source_text": plain_text})
            else:
                block = {"type": "paragraph", "text": text, "source_text": plain_text}
                if automatic_number:
                    block["automatic_number"] = automatic_number
                blocks.append(block)
            for relationship_id in paragraph._p.xpath(".//a:blip/@r:embed"):
                blocks.append({"type": "image", "relationship_id": relationship_id, "alt": "Image"})
        elif isinstance(child, CT_Tbl):
            table = Table(child, doc)
            rows = [[_table_cell_text(cell) for cell in row.cells] for row in table.rows]
            source_rows = [[_table_cell_source_text(cell) for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append({"type": "table", "rows": rows, "source_rows": source_rows})

    return blocks


def extract_embedded_images(doc: Any, document_path: str) -> dict[str, str]:
    """Write embedded document images beside the Markdown file and return paths."""
    source_path = Path(document_path).expanduser().resolve()
    assets_path = source_path.with_name(f"{source_path.stem}_assets")
    image_paths: dict[str, str] = {}
    image_number = 1
    for relationship_id, relationship in doc.part.rels.items():
        if not relationship.reltype.endswith("/image"):
            continue
        part = relationship.target_part
        extension = Path(str(part.partname)).suffix or ".bin"
        filename = f"image-{image_number}{extension}"
        assets_path.mkdir(exist_ok=True)
        (assets_path / filename).write_bytes(part.blob)
        image_paths[relationship_id] = f"{assets_path.name}/{filename}"
        image_number += 1
    return image_paths


def extract_text_blocks(content: str) -> list[dict[str, Any]]:
    """Turn any supported text-based source file into neutral content blocks."""
    return [
        {"type": "paragraph", "text": paragraph, "source_text": paragraph}
        for paragraph in re.split(r"\r?\n\s*\r?\n", content.strip())
        if paragraph.strip()
    ]


def _visible_text(text: str) -> str:
    return text.replace("**", "").replace("*", "").replace("`", "").strip()


def normalize_markdown_blocks(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Apply generic, structure-based Markdown normalization."""
    normalized = [dict(block) for block in blocks]
    for block in normalized:
        if "text" not in block:
            continue
        visible_text = _visible_text(block["text"])
        has_heading_format = "**" in block["text"]
        if block["type"] == "heading":
            if block.get("automatic_number") and not re.match(r"^\d+(?:\.\d+)*\.?\s", visible_text):
                visible_text = f"{block['automatic_number']} {visible_text}"
            block["text"] = visible_text
        elif block["type"] == "paragraph" and block.get("automatic_number") and has_heading_format:
            level = min(2 + block["automatic_number"].count("."), 6)
            block.update({"type": "heading", "level": level, "text": f"{block['automatic_number']} {visible_text}"})

    for index, block in enumerate(normalized[:-1]):
        if _visible_text(block.get("text", "")).rstrip(":").casefold() == "expected output" and normalized[index + 1]["type"] == "paragraph":
            normalized[index + 1]["type"] = "expected_output"

    in_contents = False
    for block in normalized:
        if block["type"] == "heading" and _visible_text(block.get("text", "")).rstrip(":").casefold() in {"contents", "table of contents"}:
            in_contents = True
            continue
        if in_contents and block["type"] == "ordered":
            block["toc"] = True
        elif in_contents:
            in_contents = False
    return normalized


def to_markdown(blocks: list[dict[str, Any]], image_paths: dict[str, str] | None = None) -> str:
    blocks = normalize_markdown_blocks(blocks)
    image_paths = image_paths or {}
    rendered: list[str] = []
    previous_type = ""
    for block in blocks:
        if block["type"] == "heading":
            level = block.get("level", 1)
            text = f"{'#' * level} {block['text']}"
        elif block["type"] == "bullet":
            text = f"- {block['text']}"
        elif block["type"] == "ordered":
            number = block.get("number", "1")
            label = block["text"]
            # Markdown accepts only a single integer as an ordered-list marker.
            # Use real nested lists so DOCX converters retain levels such as 2.1.
            level = number.count(".")
            list_number = number.split(".")[-1]
            text = f"{'    ' * level}{list_number}. {label}"
        elif block["type"] == "rule":
            text = "---"
        elif block["type"] == "code":
            text = f"```{block['language']}\n{block['text'].replace('`', '')}\n```"
        elif block["type"] == "expected_output":
            text = f"```text\n{block['text'].replace('`', '')}\n```"
        elif block["type"] == "image":
            image_path = image_paths.get(block["relationship_id"])
            if not image_path:
                continue
            text = f"![{block['alt']}]({image_path})"
        elif block["type"] == "table":
            rows = block["rows"]
            if rows:
                def markdown_row(row: list[str]) -> str:
                    return "| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |"

                table_lines = [markdown_row(rows[0]), "| " + " | ".join(["---"] * len(rows[0])) + " |"]
                table_lines.extend(markdown_row(row) for row in rows[1:])
                text = "\n".join(table_lines)
            else:
                continue
        else:
            text = block["text"]

        separator = "\n" if block["type"] in {"bullet", "ordered"} and previous_type == block["type"] else "\n\n"
        rendered.append(("" if not rendered else separator) + text)
        previous_type = block["type"]
    return "".join(rendered).strip() + "\n"


def to_yaml(blocks: list[dict[str, Any]]) -> str:
    """Render deterministic, valid YAML without sending conversion text to an LLM."""
    def scalar(value: str) -> str:
        # JSON double-quoted strings are valid YAML scalars and safely preserve
        # quotes, placeholders, backslashes, Unicode, and embedded newlines.
        return json.dumps(value, ensure_ascii=False)

    lines = ["content:"]
    for block in blocks:
        current_text = block.get("source_text", block.get("text", ""))
        if block["type"] == "heading":
            lines.append("  - type: heading")
            lines.append(f"    level: {block.get('level', 1)}")
            lines.append(f"    text: {scalar(current_text)}")
        elif block["type"] in {"bullet", "ordered"}:
            lines.append(f"  - type: {block['type']}")
            if block["type"] == "ordered" and "number" in block:
                lines.append(f"    number: {scalar(block['number'])}")
            lines.append(f"    text: {scalar(current_text)}")
        elif block["type"] == "rule":
            lines.append("  - type: rule")
        elif block["type"] == "code":
            lines.append("  - type: code")
            lines.append(f"    language: {scalar(block['language'])}")
            lines.append(f"    text: {scalar(current_text)}")
        elif block["type"] == "table":
            lines.append("  - type: table")
            lines.append("    rows:")
            for row in block.get("source_rows", block["rows"]):
                lines.append(f"      - {json.dumps(row, ensure_ascii=False)}")
        elif field_match := re.fullmatch(r"(.+?)\s*=\s*", current_text):
            lines.append("  - type: field")
            lines.append(f"    key: {scalar(field_match.group(1).strip())}")
            lines.append("    value: null")
        else:
            lines.append("  - type: paragraph")
            lines.append(f"    text: {scalar(current_text)}")
    return "\n".join(lines) + "\n"


def call_ollama(
    prompt: str,
    model: str = "mistral",
    base_url: str = "http://localhost:11434",
    timeout: int = 60,
    system_prompt: str | None = None,
) -> str:
    payload_data = {"model": model, "prompt": prompt, "stream": False}
    if system_prompt:
        payload_data["system"] = system_prompt
    payload = json.dumps(payload_data).encode("utf-8")
    req = request.Request(
        f"{base_url.rstrip('/')}/api/generate",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with request.urlopen(req, timeout=timeout) as response:
            data = json.load(response)
            return data.get("response", "")
    except (error.URLError, error.HTTPError, TimeoutError, json.JSONDecodeError):
        return ""


def convert_document_file(
    document_path: str,
    output_format: str = "markdown",
) -> str:
    converted = read_document_file(document_path, output_format)
    save_converted_document(document_path, output_format, converted)
    return converted


def read_document_file(
    document_path: str,
    output_format: str = "markdown",
) -> str:
    """Read any supported document into normalized text before supplying it to an LLM."""
    source = DocumentReader().read(document_path)
    blocks = extract_blocks(source.content) if source.kind == "docx" else extract_text_blocks(source.content)
    if output_format.lower() == "yaml":
        return to_yaml(blocks)
    image_paths = extract_embedded_images(source.content, document_path) if source.kind == "docx" else None
    return to_markdown(blocks, image_paths)


def converted_output_path(document_path: str, output_format: str) -> Path:
    extension = ".yaml" if output_format.lower() == "yaml" else ".md"
    source_path = Path(document_path).expanduser().resolve()
    output_path = source_path.with_suffix(extension)
    # Text inputs can already use the requested output extension. Never replace
    # the source in that case; keep the conversion beside it instead.
    if output_path == source_path:
        output_path = source_path.with_name(f"{source_path.stem}.converted{extension}")
    return output_path


def save_converted_document(document_path: str, output_format: str, content: str) -> Path:
    output_path = converted_output_path(document_path, output_format)
    output_path.write_text(content, encoding="utf-8")
    return output_path


def _words(content: str) -> set[str]:
    return {word.casefold() for word in WORD_PATTERN.findall(content)}


def assert_no_added_words(source: str, converted: str, output_format: str) -> None:
    source_counts = Counter(word.casefold() for word in WORD_PATTERN.findall(source))
    converted_counts = Counter(word.casefold() for word in WORD_PATTERN.findall(converted))
    unexpected_words = sorted((converted_counts - source_counts).keys())
    if unexpected_words:
        raise RuntimeError("Conversion contains words absent from the source: " + ", ".join(unexpected_words[:10]))
    missing_words = sorted((source_counts - converted_counts).keys())
    if missing_words:
        raise RuntimeError("Conversion is missing words from the source: " + ", ".join(missing_words[:10]))


def verify_conversion_with_ollama(
    document_path: str,
    converted: str,
    output_format: str,
    ollama_model: str = "mistral",
    base_url: str = "http://localhost:11434",
) -> bool:
    source = read_document_file(document_path, output_format)
    assert_no_added_words(source, converted, output_format)
    verdict = call_ollama(
        f"Source document:\n{source}\n\nConverted output:\n{converted}",
        model=ollama_model,
        base_url=base_url,
        system_prompt=STRICT_VALIDATION_SYSTEM_PROMPT,
    ).strip()
    verdict_words = _words(verdict)
    return "valid" in verdict_words and "invalid" not in verdict_words


async def convert_document_tool(
    document_path: str,
    output_format: str = "markdown",
    ollama_model: str = "mistral",
    base_url: str = "http://localhost:11434",
) -> str:
    converted = read_document_file(document_path, output_format)
    verify_conversion_with_ollama(document_path, converted, output_format, ollama_model, base_url)
    save_converted_document(document_path, output_format, converted)
    return converted


async def read_document_tool(
    document_path: str,
    output_format: str = "markdown",
) -> str:
    """Read a supported document into normalized Markdown or YAML for downstream LLM use."""
    return read_document_file(document_path, output_format)


async def enhance_document_tool(
    document_path: str,
    output_format: str = "markdown",
    ollama_model: str = "mistral",
    base_url: str = "http://localhost:11434",
) -> str:
    """Convert a document, then use Ollama to polish that conversion."""
    converted = read_document_file(document_path, output_format)
    prompt = (
        "Improve readability while preserving every fact from the following converted document. "
        f"Return only valid {output_format.lower()}—no commentary, code fences, or extra text.\n\n"
        f"{converted}"
    )
    enhanced = call_ollama(
        prompt,
        model=ollama_model,
        base_url=base_url,
        system_prompt=STRICT_GENERATION_SYSTEM_PROMPT,
    ).strip()
    if not enhanced:
        raise RuntimeError("Ollama did not return an enhancement. Ensure Ollama is running and the requested model is installed.")
    verify_conversion_with_ollama(document_path, enhanced, output_format, ollama_model, base_url)
    save_converted_document(document_path, output_format, enhanced + "\n")
    return enhanced + "\n"


def create_server() -> MCPServer:
    server = MCPServer(name="docflux-mcp", title="DocFlux MCP", description="Convert DOCX or text-based documents to Markdown or YAML")
    server.add_tool(
        read_document_tool,
        name="read_document",
        description="Read a supported document into normalized Markdown or YAML text for downstream LLM use",
    )
    server.add_tool(
        convert_document_tool,
        name="convert_document",
        description="Convert a document path to deterministic Markdown or YAML content",
    )
    server.add_tool(
        enhance_document_tool,
        name="enhance_document",
        description="Convert a document and polish the resulting Markdown or YAML using a local Ollama model",
    )
    return server


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the DocFlux MCP server")
    parser.add_argument("document_path", nargs="?", help="Optional supported document path to convert")
    parser.add_argument("--format", choices=["markdown", "yaml"], default="markdown", help="Output format")
    args = parser.parse_args()

    server = create_server()
    if args.document_path:
        converted = asyncio.run(convert_document_tool(args.document_path, args.format))
        print(converted)
    else:
        server.run(transport="stdio")


if __name__ == "__main__":
    main()
